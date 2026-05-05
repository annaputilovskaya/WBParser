"""Юнит-тесты для экспортёра Excel."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from pytest_mock import MockerFixture

from src.data_models.product import Product
from src.exporters.excel_exporter import ExcelExporter


class TestExcelExporter:
    """Тесты для класса ExcelExporter."""

    @pytest.fixture
    def sample_products(self) -> list[Product]:
        """Фикстура, возвращающая список тестовых товаров."""
        return [
            Product(
                url="https://example.com/1",
                article=1111111,
                name="Товар 1",
                price=5000.0,
                description="Описание 1",
                image_urls="img1.jpg",
                characteristics={"Цвет": "Красный"},
                seller_name="Продавец 1",
                seller_url="https://seller.com/1",
                sizes="44",
                stock=10,
                rating=4.8,
                review_count=100,
                country_of_origin="Россия",
                image_count=3,
            ),
            Product(
                url="https://example.com/2",
                article=2222222,
                name="Товар 2",
                price=15000.0,
                description="Описание 2",
                image_urls="img2.jpg",
                characteristics={"Материал": "Хлопок"},
                seller_name="Продавец 2",
                seller_url="https://seller.com/2",
                sizes="46,48",
                stock=5,
                rating=4.2,
                review_count=50,
                country_of_origin="Китай",
                image_count=2,
            ),
        ]

    @pytest.fixture
    def exporter(self, temp_output_dir: Path) -> ExcelExporter:
        """Фикстура, создающая экземпляр ExcelExporter с временной директорией."""
        return ExcelExporter(output_dir=temp_output_dir)

    def test_initialization_with_default_dir(self) -> None:
        """Проверяет инициализацию экспортёра с директорией по умолчанию."""
        exporter = ExcelExporter()
        assert exporter.output_dir == Path("output")

    def test_initialization_with_custom_dir(self, temp_output_dir: Path) -> None:
        """Проверяет инициализацию экспортёра с пользовательской директорией."""
        exporter = ExcelExporter(output_dir=temp_output_dir)
        assert exporter.output_dir == temp_output_dir

    def test_export_to_file_success(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет успешное сохранение файла через _export_to_file."""
        filename = "test_catalog.xlsx"
        result = exporter._export_to_file(sample_products, filename, "Тестовый каталог")
        assert result is not None
        assert result == exporter.output_dir / filename
        assert result.exists()
        assert result.stat().st_size > 0

    def test_export_to_file_empty_products(self, exporter: ExcelExporter) -> None:
        """Проверяет сохранение пустого списка товаров."""
        filename = "empty.xlsx"
        result = exporter._export_to_file([], filename, "Пустой каталог")
        assert result is not None
        assert result.exists()

    def test_export_to_file_permission_error(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        mocker: MockerFixture,
    ) -> None:
        """Проверяет обработку ошибки PermissionError при сохранении."""
        filename = "locked.xlsx"
        mocker.patch.object(
            exporter, "_save_to_excel", side_effect=PermissionError("Файл открыт")
        )
        result = exporter._export_to_file(
            sample_products, filename, "Заблокированный каталог"
        )
        assert result is None

    def test_export_to_file_general_exception(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        mocker: MockerFixture,
    ) -> None:
        """Проверяет обработку общего исключения при сохранении."""
        filename = "error.xlsx"
        mocker.patch.object(
            exporter, "_save_to_excel", side_effect=Exception("Неизвестная ошибка")
        )
        result = exporter._export_to_file(
            sample_products, filename, "Ошибочный каталог"
        )
        assert result is None

    def test_export_full_catalog(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет экспорт полного каталога."""
        result = exporter.export_full_catalog(sample_products)
        assert result is not None
        assert result.name == "full_catalog.xlsx"
        assert result.exists()

    def test_export_full_catalog_custom_filename(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет экспорт полного каталога с пользовательским именем файла."""
        custom_name = "my_catalog.xlsx"
        result = exporter.export_full_catalog(sample_products, filename=custom_name)
        assert result is not None
        assert result.name == custom_name

    def test_export_filtered_catalog_default_criteria(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет экспорт отфильтрованного каталога с критериями по умолчанию."""
        # Только первый товар соответствует критериям (рейтинг >=4.5, цена <=10000, страна Россия)
        result = exporter.export_filtered_catalog(sample_products)
        assert result is not None
        assert result.name == "filtered_catalog.xlsx"
        assert result.exists()

    @pytest.mark.parametrize(
        "min_rating, max_price, target_country, expected_count",
        [
            (4.9, 10000.0, "Россия", 0),
            (4.0, 10000.0, "Китай", 0),
            (4.0, 4000.0, "Россия", 0),
            (4.5, 10000.0, "Россия", 1),
        ],
    )
    def test_export_filtered_catalog_various_criteria(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        min_rating: float,
        max_price: float,
        target_country: str,
        expected_count: int,
    ) -> None:
        """Проверяет фильтрацию с различными критериями."""
        result = exporter.export_filtered_catalog(
            sample_products,
            min_rating=min_rating,
            max_price=max_price,
            target_country=target_country,
        )

        assert result is not None
        assert result.exists()
        if expected_count == 0:
            wb = load_workbook(result)
            ws = wb.active
            assert ws.max_row == 1  # Только заголовок

    def test_filter_products(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет метод _filter_products."""
        filtered = exporter._filter_products(
            sample_products,
            min_rating=4.5,
            max_price=10000.0,
            target_country="Россия",
        )
        assert len(filtered) == 1
        assert filtered[0].article == 1111111

    def test_filter_products_case_insensitive_country(
        self, exporter: ExcelExporter, sample_products: list[Product]
    ) -> None:
        """Проверяет регистронезависимость сравнения стран."""
        extra_product = Product(
            url="https://example.com/3",
            article=3333333,
            name="Товар 3",
            price=8000.0,
            rating=4.7,
            country_of_origin="рОсСиЯ",  # смешанный регистр
        )
        products = sample_products + [extra_product]
        filtered = exporter._filter_products(
            products,
            min_rating=4.0,
            max_price=20000.0,
            target_country="РОССИЯ",  # верхний регистр
        )
        assert len(filtered) == 2
        articles = {p.article for p in filtered}
        assert 1111111 in articles
        assert 3333333 in articles

    def test_filter_products_empty_country_excluded(
        self, exporter: ExcelExporter
    ) -> None:
        """Проверяет исключение товаров с пустой страной производства."""
        product_without_country = Product(
            url="https://example.com/4",
            article=4444444,
            name="Товар 4",
            price=7000.0,
            rating=4.9,
            country_of_origin="",  # пустая строка
        )
        product_with_none_country = Product(
            url="https://example.com/5",
            article=5555555,
            name="Товар 5",
            price=6000.0,
            rating=4.8,
            country_of_origin=None,  # None
        )
        products = [product_without_country, product_with_none_country]
        filtered = exporter._filter_products(
            products,
            min_rating=4.0,
            max_price=10000.0,
            target_country="Россия",
        )
        assert len(filtered) == 0

    def test_save_to_excel_creates_file(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        temp_output_dir: Path,
    ) -> None:
        """Проверяет, что _save_to_excel создаёт корректный Excel файл."""
        rows = [p.to_dict() for p in sample_products]
        filepath = temp_output_dir / "test_save.xlsx"
        exporter._save_to_excel(rows, filepath)
        assert filepath.exists()

    def test_save_to_excel_headers_mapping(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        temp_output_dir: Path,
    ) -> None:
        """Проверяет соответствие заголовков колонок."""
        rows = [sample_products[0].to_dict()]
        filepath = temp_output_dir / "headers_test.xlsx"
        exporter._save_to_excel(rows, filepath)
        # Загружаем файл и проверяем первую строку
        wb = load_workbook(filepath)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        expected_headers = list(exporter.HEADERS_MAP.values())
        assert headers == expected_headers

    def test_save_to_excel_data_rows(
        self,
        exporter: ExcelExporter,
        sample_products: list[Product],
        temp_output_dir: Path,
    ) -> None:
        """Проверяет корректность данных в строках."""
        rows = [sample_products[0].to_dict()]
        filepath = temp_output_dir / "data_test.xlsx"
        exporter._save_to_excel(rows, filepath)

        wb = load_workbook(filepath)
        ws = wb.active
        # Вторая строка должна содержать данные
        row_data = [cell.value for cell in ws[2]]
        expected_values = [
            sample_products[0].url,
            sample_products[0].article,
            sample_products[0].name,
            sample_products[0].price,
            sample_products[0].description,
            sample_products[0].image_urls,
            "Цвет: Красный",
            sample_products[0].seller_name,
            sample_products[0].seller_url,
            sample_products[0].sizes,
            sample_products[0].stock,
            sample_products[0].rating,
            sample_products[0].review_count,
        ]
        assert row_data == expected_values
